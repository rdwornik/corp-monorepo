"""Tests for light_scan module — format scanners and ScanResult contract."""
from unittest.mock import MagicMock, patch

from corp.ingest.light_scan import ScanResult, light_scan

# ---------------------------------------------------------------------------
# ScanResult contract
# ---------------------------------------------------------------------------


def test_scan_result_filename_text():
    """filename_text returns just the filename."""
    r = ScanResult(filename="report.pdf", extension=".pdf", title="Report Title")
    assert r.filename_text == "report.pdf"


def test_scan_result_content_text_combines_fields():
    """content_text concatenates title + headers + snippet."""
    r = ScanResult(
        filename="f.docx",
        extension=".docx",
        title="My Title",
        headers=["Section A", "Section B"],
        snippet="Body text here",
    )
    text = r.content_text
    assert "My Title" in text
    assert "Section A" in text
    assert "Body text here" in text


def test_scan_result_separate_features():
    """filename_text and content_text are independent (Council: separate spaces)."""
    r = ScanResult(
        filename="andy.pdf",
        extension=".pdf",
        title="Demand Planning Training Q3",
        snippet="quarterly review content",
    )
    # Content has training signal; filename has none
    assert "Demand Planning" in r.content_text
    assert "Demand Planning" not in r.filename_text
    # Filename has file name; content does not include it
    assert "andy.pdf" in r.filename_text
    assert "andy.pdf" not in r.content_text


def test_scan_result_content_text_empty_fields():
    """content_text is empty string when all content fields are empty."""
    r = ScanResult(filename="empty.pdf", extension=".pdf")
    assert r.content_text == ""


def test_feature_dict_has_all_keys():
    """to_feature_dict returns expected structure."""
    r = ScanResult(
        filename="deck.pptx",
        extension=".pptx",
        title="Sales Deck",
        headers=["Intro", "Demo"],
        slide_count=10,
        scan_tier="full",
    )
    d = r.to_feature_dict()
    expected_keys = {
        "filename",
        "extension",
        "title",
        "header_count",
        "slide_count",
        "page_count",
        "duration_seconds",
        "sheet_count",
        "word_count",
        "content_text_length",
        "scan_tier",
        "has_title",
        "has_agenda",
    }
    assert expected_keys == set(d.keys())
    assert d["has_title"] is True
    assert d["header_count"] == 2


def test_feature_dict_has_agenda_detection():
    """has_agenda is True when headers contain 'agenda'."""
    r = ScanResult(
        filename="meeting.docx",
        extension=".docx",
        headers=["Welcome", "Today's Agenda", "Next Steps"],
    )
    assert r.to_feature_dict()["has_agenda"] is True


def test_feature_dict_no_agenda():
    """has_agenda is False when no header contains 'agenda'."""
    r = ScanResult(filename="f.pdf", extension=".pdf", headers=["Introduction", "Summary"])
    assert r.to_feature_dict()["has_agenda"] is False


# ---------------------------------------------------------------------------
# PPTX scanner
# ---------------------------------------------------------------------------


def test_pptx_extracts_slide_titles(tmp_path):
    """PPTX scan returns slide titles and first-slide snippet."""
    pptx_path = tmp_path / "deck.pptx"
    prs = _make_pptx(["Welcome to Training", "Agenda", "Module 1"])
    prs.save(str(pptx_path))

    result = light_scan(pptx_path)

    assert result.scan_tier == "full"
    assert result.title == "Welcome to Training"
    assert "Agenda" in result.slide_titles
    assert result.slide_count == 3


def test_pptx_missing_library_graceful(tmp_path):
    """Missing python-pptx returns error, doesn't crash."""
    pptx_path = tmp_path / "deck.pptx"
    pptx_path.write_bytes(b"fake")

    with patch.dict("sys.modules", {"pptx": None}):
        result = light_scan(pptx_path)

    assert len(result.errors) > 0
    assert result.scan_tier in ("degraded", "filename_only")


# ---------------------------------------------------------------------------
# DOCX scanner
# ---------------------------------------------------------------------------


def test_docx_extracts_headings(tmp_path):
    """DOCX scan returns Heading 1/2 as headers."""
    from docx import Document

    doc_path = tmp_path / "report.docx"
    doc = Document()
    doc.add_heading("Project Overview", level=1)
    doc.add_heading("Background", level=2)
    doc.add_paragraph("This is the body text of the document.")
    doc.save(str(doc_path))

    result = light_scan(doc_path)

    assert result.scan_tier == "full"
    assert result.title == "Project Overview"
    assert "Background" in result.headers
    assert "body text" in result.snippet


def test_docx_missing_library_graceful(tmp_path):
    """Missing python-docx returns error, doesn't crash."""
    doc_path = tmp_path / "report.docx"
    doc_path.write_bytes(b"fake")

    with patch.dict("sys.modules", {"docx": None}):
        result = light_scan(doc_path)

    assert len(result.errors) > 0
    assert result.scan_tier in ("degraded", "filename_only")


# ---------------------------------------------------------------------------
# PDF scanner
# ---------------------------------------------------------------------------


def test_pdf_extracts_first_page(tmp_path):
    """PDF scan returns title from first line and snippet."""

    pdf_path = tmp_path / "doc.pdf"

    # Create a real minimal PDF using pdfplumber's dependency (pdfminer via reportlab or fpdf)
    # Use fpdf2 or just mock pdfplumber since we don't want heavy test deps
    mock_page = MagicMock()
    mock_page.extract_text.return_value = (
        "Demand Planning Training Q3 2025\nSection 1\nSection 2\nBody text."
    )

    mock_pdf = MagicMock()
    mock_pdf.__enter__ = lambda s: s
    mock_pdf.__exit__ = MagicMock(return_value=False)
    mock_pdf.pages = [mock_page]

    pdf_path.write_bytes(b"%PDF-1.4 fake")

    with patch("pdfplumber.open", return_value=mock_pdf):
        result = light_scan(pdf_path)

    assert result.title == "Demand Planning Training Q3 2025"
    assert "Section 1" in result.headers
    assert result.scan_tier == "full"


# ---------------------------------------------------------------------------
# XLSX scanner
# ---------------------------------------------------------------------------


def test_xlsx_extracts_sheet_names(tmp_path):
    """XLSX scan returns sheet names and column headers."""
    from openpyxl import Workbook

    xlsx_path = tmp_path / "data.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales Data"
    ws.append(["Date", "Product", "Revenue", "Units"])
    ws.append(["2025-01-01", "WMS", 50000, 10])
    wb.save(str(xlsx_path))

    result = light_scan(xlsx_path)

    assert result.scan_tier == "full"
    assert "Sales Data" in result.sheet_names
    assert "Date" in result.column_headers
    assert "Product" in result.column_headers


# ---------------------------------------------------------------------------
# CSV scanner
# ---------------------------------------------------------------------------


def test_csv_extracts_headers(tmp_path):
    """CSV scan returns column headers + first rows as snippet."""
    csv_path = tmp_path / "items.csv"
    csv_path.write_text("item_id,description,price\n1,Widget,9.99\n2,Gadget,19.99\n")

    result = light_scan(csv_path)

    assert result.scan_tier == "full"
    assert "item_id" in result.column_headers
    assert "description" in result.column_headers
    assert "item_id" in result.title  # title built from headers


def test_csv_empty_file_graceful(tmp_path):
    """Empty CSV doesn't crash."""
    csv_path = tmp_path / "empty.csv"
    csv_path.write_text("")

    result = light_scan(csv_path)
    assert result.scan_tier in ("full", "degraded", "filename_only")
    assert result.filename == "empty.csv"


# ---------------------------------------------------------------------------
# TXT/MD scanner
# ---------------------------------------------------------------------------


def test_text_extracts_title_and_headers(tmp_path):
    """TXT/MD scan returns first line as title and markdown headers."""
    md_path = tmp_path / "notes.md"
    md_path.write_text(
        "# Meeting Notes\n\n## Agenda\n\nSome body text here.\n\n## Action Items\n"
    )

    result = light_scan(md_path)

    assert result.scan_tier == "full"
    assert result.title == "Meeting Notes"
    assert "Agenda" in result.headers
    assert "Action Items" in result.headers


def test_txt_plain_text(tmp_path):
    """Plain .txt file works without markdown headers."""
    txt_path = tmp_path / "readme.txt"
    txt_path.write_text("Training guide for new starters.\nLine two.\n")

    result = light_scan(txt_path)

    assert result.scan_tier == "full"
    assert result.title == "Training guide for new starters."


# ---------------------------------------------------------------------------
# MP4 scanner
# ---------------------------------------------------------------------------


def test_mp4_extracts_duration_from_ffprobe(tmp_path):
    """MP4 scan returns duration when ffprobe is available."""
    mp4_path = tmp_path / "recording.mp4"
    mp4_path.write_bytes(b"fake video")

    ffprobe_output = '{"format": {"duration": "3600.0", "tags": {"title": "Q3 Review Meeting"}}}'

    mock_proc = MagicMock()
    mock_proc.returncode = 0
    mock_proc.stdout = ffprobe_output

    with patch("subprocess.run", return_value=mock_proc):
        result = light_scan(mp4_path)

    assert result.duration_seconds == 3600.0
    assert result.title == "Q3 Review Meeting"
    assert result.scan_tier == "full"


def test_mp4_long_duration_hint(tmp_path):
    """Long MP4 (>45min) adds 'long recording likely meeting' hint."""
    mp4_path = tmp_path / "webinar.mp4"
    mp4_path.write_bytes(b"fake")

    ffprobe_output = '{"format": {"duration": "3600.0", "tags": {}}}'
    mock_proc = MagicMock()
    mock_proc.returncode = 0
    mock_proc.stdout = ffprobe_output

    with patch("subprocess.run", return_value=mock_proc):
        result = light_scan(mp4_path)

    assert "long recording likely meeting" in result.snippet


def test_mp4_short_duration_hint(tmp_path):
    """Short MP4 (<5min) adds 'short clip likely demo' hint."""
    mp4_path = tmp_path / "demo_clip.mp4"
    mp4_path.write_bytes(b"fake")

    ffprobe_output = '{"format": {"duration": "180.0", "tags": {}}}'
    mock_proc = MagicMock()
    mock_proc.returncode = 0
    mock_proc.stdout = ffprobe_output

    with patch("subprocess.run", return_value=mock_proc):
        result = light_scan(mp4_path)

    assert "short clip likely demo" in result.snippet


def test_mp4_ffprobe_not_found(tmp_path):
    """Missing ffprobe returns error, doesn't crash."""
    mp4_path = tmp_path / "video.mp4"
    mp4_path.write_bytes(b"fake")

    with patch("subprocess.run", side_effect=FileNotFoundError("ffprobe not found")):
        result = light_scan(mp4_path)

    assert any("ffprobe not found" in e for e in result.errors)
    assert result.filename == "video.mp4"


def test_mp4_ffprobe_timeout(tmp_path):
    """ffprobe timeout degrades gracefully."""
    import subprocess

    mp4_path = tmp_path / "huge.mp4"
    mp4_path.write_bytes(b"fake")

    with patch("subprocess.run", side_effect=subprocess.TimeoutExpired("ffprobe", 3)):
        result = light_scan(mp4_path)

    assert any("timeout" in e.lower() for e in result.errors)


# ---------------------------------------------------------------------------
# Fault tolerance
# ---------------------------------------------------------------------------


def test_corrupted_pptx_degrades_gracefully(tmp_path):
    """Corrupted PPTX degrades to filename_only, doesn't crash."""
    bad = tmp_path / "corrupt.pptx"
    bad.write_bytes(b"not a real pptx file at all")

    result = light_scan(bad)

    assert result.scan_tier in ("degraded", "filename_only")
    assert len(result.errors) > 0
    assert result.filename == "corrupt.pptx"


def test_unsupported_extension(tmp_path):
    """Unsupported file extension returns filename_only with error."""
    f = tmp_path / "file.xyz"
    f.write_bytes(b"content")

    result = light_scan(f)

    assert result.scan_tier == "filename_only"
    assert any("Unsupported" in e for e in result.errors)


def test_scan_time_recorded(tmp_path):
    """scan_time_ms is populated after scan."""
    txt = tmp_path / "f.txt"
    txt.write_text("hello world")

    result = light_scan(txt)

    assert result.scan_time_ms > 0


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _make_pptx(slide_titles: list[str]):
    """Create a minimal PPTX with titled slides."""
    from pptx import Presentation

    prs = Presentation()
    blank_layout = prs.slide_layouts[0]  # title slide layout

    for title_text in slide_titles:
        slide = prs.slides.add_slide(blank_layout)
        title = slide.shapes.title
        if title:
            title.text = title_text

    return prs
